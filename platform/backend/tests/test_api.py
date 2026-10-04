"""认证与业务接口测试（农户幂等提交、收购单流转、OCR 兼容契约、Excel 导出）。"""

import base64
from datetime import date
from io import BytesIO

from openpyxl import load_workbook

from app.services import bailian


def test_health_contract(client):
    response = client.get("/api/health")
    assert response.status_code == 200
    body = response.json()
    assert body["ok"] is True
    assert body["provider"] == "bailian"
    assert "bailianKeyReady" in body


def test_ocr_contract_and_mocked_bailian(client, monkeypatch):
    async def fake_recognize(image_bytes, extra_hint="", mime="image/jpeg"):
        assert image_bytes
        return {
            "name": "张三",
            "idNumber": "11010519491231002X",
            "address": "北京市海淀区中关村大街27号",
            "confidence": {"name": 0.95, "idNumber": 0.98, "address": 0.9},
        }

    async def fake_address(image_bytes, extra_hint="", mime="image/jpeg"):
        return {"address": "北京市海淀区中关村大街27号", "confidence": 0.9}

    monkeypatch.setattr(bailian, "recognize", fake_recognize)
    monkeypatch.setattr(bailian, "recognize_address", fake_address)


    from PIL import Image as PILImage
    import io as _io
    tiny = _io.BytesIO()
    PILImage.new("RGB", (64, 48), (220, 220, 220)).save(tiny, format="JPEG", quality=85)
    png = base64.b64encode(tiny.getvalue()).decode()
    response = client.post("/api/ocr/idcard", json={"imageBase64": png, "mime": "image/jpeg"})
    assert response.status_code == 200
    body = response.json()
    assert body["ok"] is True
    assert body["provider"] == "bailian"
    assert body["name"] == "张三"
    assert body["idNumber"] == "11010519491231002X"
    assert body["address"] == "北京市海淀区中关村大街27号"
    assert body["warnings"] == []
    assert body["checksumOk"] is True
    assert body["addressVerified"] is True  # 两角度读数一致
    assert body["confidence"]["idNumber"] == 0.98

    raw = client.post(
        "/api/ocr/idcard", content=tiny.getvalue(), headers={"Content-Type": "image/jpeg"}
    )
    assert raw.status_code == 200
    assert raw.json()["idNumber"] == "11010519491231002X"


def test_ocr_rejects_bad_payloads(client):
    assert client.post("/api/ocr/idcard", json={}).status_code == 400
    big = b"x" * (10 * 1024 * 1024 + 1)
    response = client.post(
        "/api/ocr/idcard", content=big, headers={"Content-Type": "image/jpeg"}
    )
    assert response.status_code == 413


def test_auth_flow_and_protection(client):
    assert client.get("/api/auth/status").json()["needsSetup"] is True
    assert client.get("/api/farmers").status_code == 401

    setup = client.post("/api/auth/setup", json={"password": "grain123"})
    assert setup.status_code == 200
    token = setup.json()["token"]

    assert client.post("/api/auth/setup", json={"password": "again123"}).status_code == 409
    assert client.post("/api/auth/login", json={"password": "wrong"}).status_code == 401

    login = client.post("/api/auth/login", json={"password": "grain123"})
    assert login.status_code == 200 and login.json()["token"]

    client.headers.update({"Authorization": f"Bearer {token}"})
    assert client.get("/api/farmers").status_code == 200


def test_farmer_submit_idempotent_and_records_flow(client):
    setup = client.post("/api/auth/setup", json={"password": "grain123"})
    client.headers.update({"Authorization": f"Bearer {setup.json()['token']}"})

    payload = {
        "name": "李四",
        "idNumber": "11010519491231002X",
        "address": "北京市海淀区中关村大街27号",
    }
    first = client.post("/api/farmers/submit", json=payload)
    assert first.status_code == 200 and first.json()["created"] is True
    farmer_id = first.json()["farmer"]["id"]

    second = client.post("/api/farmers/submit", json={**payload, "phone": "13800000000"})
    assert second.status_code == 200 and second.json()["created"] is False
    assert second.json()["farmer"]["phone"] == "13800000000"

    listed = client.get("/api/farmers", params={"q": "李四"})
    assert listed.status_code == 200
    body = listed.json()
    assert body["total"] == 1 and body["items"][0]["source"] == "miniprogram"

    record = {
        "farmerId": farmer_id,
        "date": date.today().isoformat(),
        "product": "玉米",
        "vehicleNo": "京A12345",
        "grade": "2",
        "moisture": 16.0,
        "impurity": 0.0,
        "moldy": 0.0,
        "gross": 4918,
        "tare": 1858,
        "deductSite": 0,
        "deductMoisture": 130,
        "deductImpurity": 0,
        "deductMoldy": 0,
        "pricePerKg": 2.36,
        "unloadFee": 0,
        "warehouse": "1号仓",
        "weigher": "王检查",
    }
    created = client.post("/api/records", json=record)
    assert created.status_code == 200
    record_body = created.json()
    assert record_body["receiptNo"] == date.today().strftime("%Y%m%d") + "001"
    assert record_body["pricePerKg"] == 2.36
    assert record_body["netWeight"] == 3060
    assert record_body["pricedWeight"] == 2930
    assert record_body["totalAmount"] == 6914.8
    assert record_body["status"] == "draft"

    record_id = record_body["id"]
    updated = client.put(f"/api/records/{record_id}", json={"pricePerKg": 2.4})
    assert updated.status_code == 200
    assert updated.json()["totalAmount"] == 7032.0

    confirmed = client.post(f"/api/records/{record_id}/confirm")
    assert confirmed.status_code == 200 and confirmed.json()["status"] == "confirmed"

    locked = client.put(f"/api/records/{record_id}", json={"pricePerKg": 2.5})
    assert locked.status_code == 409

    listed = client.get("/api/records", params={"status": "confirmed"})
    assert listed.json()["total"] == 1
    stats = client.get("/api/stats/summary").json()
    assert stats["today"]["count"] == 1
    assert stats["today"]["tons"] == 2.93
    assert stats["today"]["amount"] == 7032.0


def test_export_excel(auth_client):
    payload = {
        "name": "李四",
        "idNumber": "11010519491231002X",
        "address": "北京市海淀区中关村大街27号",
    }
    farmer_id = auth_client.post("/api/farmers/submit", json=payload).json()["farmer"]["id"]
    today = date.today().isoformat()

    first = auth_client.post(
        "/api/records",
        json={
            "farmerId": farmer_id,
            "date": today,
            "product": "玉米",
            "gross": 4918,
            "tare": 1858,
            "deductMoisture": 130,
            "pricePerKg": 2.36,
        },
    )
    assert first.status_code == 200
    auth_client.post(f"/api/records/{first.json()['id']}/confirm")
    second = auth_client.post(
        "/api/records",
        json={"farmerId": farmer_id, "date": today, "product": "玉米", "gross": 1000, "tare": 400, "pricePerKg": 1.0},
    )
    assert second.status_code == 200

    response = auth_client.get("/api/records/export.xlsx")
    assert response.status_code == 200
    assert "spreadsheetml" in response.headers["content-type"]

    workbook = load_workbook(BytesIO(response.content))
    assert workbook.sheetnames == ["收购明细", "汇总统计"]

    sheet = workbook["收购明细"]
    assert sheet.cell(row=1, column=1).value == "编号"
    assert sheet.cell(row=2, column=1).value == first.json()["receiptNo"]
    assert sheet.cell(row=3, column=1).value == second.json()["receiptNo"]
    assert sheet.cell(row=4, column=1).value == "合计（不含作废，共 2 张）"
    assert abs(sheet.cell(row=4, column=26).value - 7514.8) < 0.001

    summary = workbook["汇总统计"]
    assert summary["A1"].value == "收购统计汇总"
    assert summary["A4"].value == "单据总数"
    assert summary["B4"].value == 2
