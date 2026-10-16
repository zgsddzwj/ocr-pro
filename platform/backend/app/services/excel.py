"""收购单导出 Excel：「收购明细」+「汇总统计」两个工作表。"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..models import Farmer, PurchaseRecord

STATUS_TEXT = {'draft': '草稿', 'confirmed': '已确认', 'void': '已作废'}

DETAIL_HEADERS = [
    '编号', '日期', '状态', '农户姓名', '身份证号', '地址', '产品名称', '产品等级',
    '水分%', '杂质霉变%', '毛重(公斤)', '皮重(公斤)', '综合扣量(公斤)', '净重(公斤)',
    '单价(元/公斤)', '市斤价', '金额', '金额大写',
    '卸车货位', '银行账号', '联系电话', '化验员', '监管员', '进出时间',
]

# 1-based 列号 -> 数字格式
DETAIL_FORMATS = {
    9: '0.00', 10: '0.00',
    11: '#,##0.000', 12: '#,##0.000', 13: '#,##0.000', 14: '#,##0.000',
    15: '0.0000', 16: '0.0000',
    17: '#,##0.00',
}

DETAIL_WIDTHS = {
    1: 14, 2: 11, 3: 9, 4: 10, 5: 20, 6: 30, 7: 10, 8: 9,
    9: 8, 10: 11, 11: 12, 12: 12, 13: 14, 14: 12,
    15: 13, 16: 10, 17: 12, 18: 26,
    19: 10, 20: 20, 21: 13, 22: 8, 23: 8, 24: 16,
}

# 合计行需要求和的列（毛重/皮重/综合扣量/净重/金额）
SUM_COLUMNS = [11, 12, 13, 14, 17]

DETAIL_FIELDS = {
    11: 'gross',
    12: 'tare',
    13: 'deduct_total',
    14: 'net_weight',
    17: 'grain_amount',
}

HEADER_FILL = PatternFill('solid', fgColor='1E6F43')
HEADER_FONT = Font(color='FFFFFF', bold=True)
TOTAL_FONT = Font(bold=True)


def _in_out_text(record: PurchaseRecord) -> str:
    if not record.weigh_in_at:
        return ''
    return f'{record.weigh_in_at}{" / " + record.weigh_out_at if record.weigh_out_at else ""}'


def _detail_row(record: PurchaseRecord, farmer: Farmer | None) -> list:
    return [
        record.receipt_no,
        record.date,
        STATUS_TEXT.get(record.status, record.status),
        farmer.name if farmer else '',
        farmer.id_number if farmer else '',
        farmer.address if farmer else '',
        record.product,
        record.grade,
        record.moisture,
        record.impurity,
        record.gross,
        record.tare,
        record.deduct_total,
        record.net_weight,
        record.price_per_kg,
        record.price_per_jin,
        record.grain_amount,
        record.amount_cn,
        record.unload_spot,
        record.bank_card,
        record.phone,
        record.checker,
        record.supervisor,
        _in_out_text(record),
    ]


def _summarize(records: list[PurchaseRecord]) -> dict:
    """非作废口径的合计。"""
    active = [record for record in records if record.status != 'void']
    weight = sum(record.net_weight for record in active)
    return {
        'count': len(active),
        'gross': sum(record.gross for record in active),
        'net': weight,
        'tons': weight / 1000,
        'amount': sum(record.grain_amount for record in active),
    }


def _style_header(sheet, ncols: int) -> None:
    for col in range(1, ncols + 1):
        cell = sheet.cell(row=1, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal='center', vertical='center')
    sheet.freeze_panes = 'A2'


def _apply_formats(sheet, formats: dict, first_row: int, last_row: int) -> None:
    for row in range(first_row, last_row + 1):
        for col, fmt in formats.items():
            sheet.cell(row=row, column=col).number_format = fmt


def _set_widths(sheet, widths: dict) -> None:
    for col, width in widths.items():
        sheet.column_dimensions[get_column_letter(col)].width = width


def build_workbook(pairs: list[tuple[PurchaseRecord, Farmer | None]]) -> bytes:
    """pairs: (收购单, 农户)，按日期+编号排序后的全集。"""
    buffer = BytesIO()
    workbook = Workbook()

    # ---------- Sheet1 收购明细 ----------
    detail = workbook.active
    detail.title = '收购明细'
    detail.append(DETAIL_HEADERS)
    _style_header(detail, len(DETAIL_HEADERS))

    for record, farmer in pairs:
        detail.append(_detail_row(record, farmer))

    if pairs:
        active = [record for record, _farmer in pairs if record.status != 'void']
        total_row = detail.max_row + 1
        detail.cell(row=total_row, column=1, value=f'合计（不含作废，共 {len(active)} 张）')
        for col in SUM_COLUMNS:
            value = sum(getattr(record, DETAIL_FIELDS[col]) for record in active)
            cell = detail.cell(row=total_row, column=col, value=round(value, 3))
            cell.number_format = DETAIL_FORMATS.get(col, '#,##0.00')
        for col in range(1, len(DETAIL_HEADERS) + 1):
            detail.cell(row=total_row, column=col).font = TOTAL_FONT
        _apply_formats(detail, DETAIL_FORMATS, 2, total_row - 1)
    _set_widths(detail, DETAIL_WIDTHS)

    # ---------- Sheet2 汇总统计 ----------
    summary = workbook.create_sheet('汇总统计')
    for col, width in zip('ABCD', (24, 16, 18, 18)):
        summary.column_dimensions[col].width = width

    bold = Font(bold=True)
    title = Font(bold=True, size=14)

    summary['A1'] = '收购统计汇总'
    summary['A1'].font = title

    counts = {'draft': 0, 'confirmed': 0, 'void': 0}
    for record, _farmer in pairs:
        if record.status in counts:
            counts[record.status] += 1
    overview = _summarize([record for record, _farmer in pairs])

    row = 3
    summary.cell(row=row, column=1, value='总览').font = bold
    for label, value in [
        ('单据总数', len(pairs)),
        ('已确认', counts['confirmed']),
        ('草稿', counts['draft']),
        ('已作废', counts['void']),
        ('车次（不含作废）', overview['count']),
        ('毛重合计（公斤）', round(overview['gross'], 3)),
        ('净重合计（公斤）', round(overview['net'], 3)),
        ('净重合计（吨）', round(overview['tons'], 3)),
        ('金额合计（元）', round(overview['amount'], 2)),
    ]:
        row += 1
        summary.cell(row=row, column=1, value=label)
        summary.cell(row=row, column=2, value=value)

    by_date: dict[str, list[PurchaseRecord]] = {}
    by_product: dict[str, list[PurchaseRecord]] = {}
    for record, _farmer in pairs:
        if record.status == 'void':
            continue
        by_date.setdefault(record.date, []).append(record)
        by_product.setdefault(record.product or '（未填）', []).append(record)

    def write_group(sheet_row: int, heading: str, groups: dict, label_header: str) -> int:
        summary.cell(row=sheet_row, column=1, value=heading).font = bold
        sheet_row += 1
        for col, header in enumerate([label_header, '车数', '净重(吨)', '金额(元)'], start=1):
            summary.cell(row=sheet_row, column=col, value=header).font = bold
        for name in sorted(groups):
            stats = _summarize(groups[name])
            sheet_row += 1
            summary.cell(row=sheet_row, column=1, value=name)
            summary.cell(row=sheet_row, column=2, value=stats['count'])
            summary.cell(row=sheet_row, column=3, value=round(stats['tons'], 3))
            summary.cell(row=sheet_row, column=4, value=round(stats['amount'], 2))
        return sheet_row

    row = write_group(row + 2, '按日期汇总（不含作废）', by_date, '日期')
    write_group(row + 2, '按品名汇总（不含作废）', by_product, '品名')

    workbook.save(buffer)
    return buffer.getvalue()
